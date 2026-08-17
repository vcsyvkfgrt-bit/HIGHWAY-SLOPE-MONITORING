"""
边坡监测 Excel 累计值提取脚本（Python 版）

功能：
1. 读取多个边坡原始 Excel 文件。
2. 从“成果统计表”中提取：
   - 地表偏移 / 水平位移：累计偏移(mm)、累计位移(mm)
   - 沉降数据：累计沉降(mm)
3. 分别输出为系统录入模板格式：
   监测日期 | 测点1 | 测点2 | ... | 备注

运行：
    python tools/extract_slope_cumulative_to_template.py

依赖：
    pip install openpyxl

说明：
- 当前脚本优先读取“成果统计表”，因为该表已经是整理后的累计成果。
- 如果后续表格样式变化，优先修改 CONFIG、TYPES、extract_from_summary_sheet()。
"""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
import re
import sys

try:
    from openpyxl import load_workbook, Workbook
except ImportError:
    print("缺少依赖 openpyxl，请先执行：pip install openpyxl")
    sys.exit(1)


CONFIG = {
    "input_files": [
        r"E:/汪斌-资料/兴长项目/项目部季度考核材料/监测相关/原始监测数据/三标/地表水平位移/DK0+206.5-EK1+545.5.xlsx",
        r"E:/汪斌-资料/兴长项目/项目部季度考核材料/监测相关/原始监测数据/三标/地表水平位移/H10.xlsx",
        r"E:/汪斌-资料/兴长项目/项目部季度考核材料/监测相关/原始监测数据/三标/地表水平位移/H24.xlsx",
        r"E:/汪斌-资料/兴长项目/项目部季度考核材料/监测相关/原始监测数据/三标/地表水平位移/三星店大桥终点桥台ZK72+210-284沉降板.xlsx",
    ],
    "output_dir": r"E:/code/周报月报平台/converted_slope_cumulative_templates_python",
    "summary_sheet_name": "成果统计表",
}


TYPES = {
    "surface": {
        "label": "地表偏移数据",
        "output_sheet_name": "地表位移监测点",
        "title_keywords": ["累计偏移成果数据表", "累计位移成果数据表"],
        "value_keywords": ["累计偏移", "累计位移"],
    },
    "settlement": {
        "label": "沉降数据",
        "output_sheet_name": "沉降监测点",
        "title_keywords": ["累计沉降成果数据表"],
        "value_keywords": ["累计沉降"],
    },
}


SLOPE_SPLIT_RULES = [
    {
        "source_base_name": "DK0+206.5-EK1+545.5",
        "target_base_name": "EK1+335~EK1+435",
        "points": {"DQ-1", "DQ-2", "DQ-3"},
    },
]


def normalize_text(value) -> str:
    """统一处理单元格文本。"""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def safe_name(name: str) -> str:
    """生成安全文件名。"""
    name = normalize_text(name)
    name = re.sub(r'[\\/:*?"<>|]', "_", name)
    return name[:100] or "converted"


def parse_number(value):
    """解析数值，失败返回 None。"""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 4)
    text = normalize_text(value).replace(",", "")
    try:
        return round(float(text), 4)
    except ValueError:
        return None


def parse_date(value) -> str:
    """
    解析日期，统一输出 YYYY-MM-DD。

    兼容：
    - datetime
    - Excel 日期
    - 5/7/25
    - 2025/7/9 9:30
    - 2025-07-09
    - 2025年7月9日
    """
    if value is None or value == "":
        return ""

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    text = normalize_text(value)
    if not text:
        return ""

    date_part = text.split()[0]
    date_part = (
        date_part.replace("年", "/")
        .replace("月", "/")
        .replace("日", "")
        .replace("号", "")
        .replace(".", "/")
        .replace("-", "/")
    )

    match = re.match(r"^(\d{4})/(\d{1,2})/(\d{1,2})$", date_part)
    if match:
        y, m, d = match.groups()
        return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"

    match = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{2,4})$", date_part)
    if match:
        m, d, y = match.groups()
        year = int(y)
        if year < 100:
            year += 2000
        return f"{year:04d}-{int(m):02d}-{int(d):02d}"

    return ""


def normalize_point_name(name: str) -> str:
    """规范测点名称。"""
    text = normalize_text(name)
    text = re.sub(r"^新", "", text)
    text = text.replace("#", "")
    return text


def normalize_point_key(name: str) -> str:
    return re.sub(r"\s+", "", normalize_point_name(name)).upper()


def includes_any(text, keywords) -> bool:
    text = normalize_text(text)
    return any(keyword in text for keyword in keywords)


def sheet_to_rows(ws):
    """把工作表转成二维数组。"""
    return [
        [cell for cell in row]
        for row in ws.iter_rows(values_only=True)
    ]


def cell(rows, row_index: int, col_index: int):
    """安全读取二维数组单元格。"""
    if row_index < 0 or row_index >= len(rows):
        return None
    row = rows[row_index]
    if col_index < 0 or col_index >= len(row):
        return None
    return row[col_index]


def max_cols(rows) -> int:
    return max((len(row) for row in rows), default=0)


def find_block_starts(rows, type_config):
    """
    查找成果统计表中的数据块起点。

    一般第一行标题类似：
    - xxx测点累计偏移成果数据表
    - xxx测点累计沉降成果数据表
    """
    starts = []
    scan_rows = min(len(rows), 10)

    for r in range(scan_rows):
        for c, value in enumerate(rows[r]):
            if includes_any(value, type_config["title_keywords"]):
                starts.append({"title_row": r, "start_col": c})

    return starts


def find_next_block_start(block_starts, current_start_col: int, fallback_end_col: int) -> int:
    later_cols = sorted(
        item["start_col"]
        for item in block_starts
        if item["start_col"] > current_start_col
    )
    return later_cols[0] if later_cols else fallback_end_col


def build_point_columns(rows, start_col: int, end_col: int, type_config):
    """
    根据“累计偏移 / 累计沉降”列，反推对应测点列。

    成果统计表常见结构：
    第 2 行：点号、BP-1、空、BP-2、空...
    第 4 行：空、速率、累计值、速率、累计值...

    因此遇到“累计值”列时，向左寻找最近的测点名称。
    """
    point_row_index = 1
    header_row_index = 3
    point_columns = []

    for c in range(start_col + 1, end_col):
        header = normalize_text(cell(rows, header_row_index, c))
        if not includes_any(header, type_config["value_keywords"]):
            continue

        point_name = ""
        for pc in range(c, start_col - 1, -1):
            point_name = normalize_point_name(cell(rows, point_row_index, pc))
            if point_name:
                break

        if point_name:
            point_columns.append({
                "point": point_name,
                "col": c,
            })

    return point_columns


def extract_block_records(rows, block, all_blocks, type_key: str):
    """提取一个成果统计数据块中的累计值。"""
    type_config = TYPES[type_key]
    end_col = find_next_block_start(all_blocks, block["start_col"], max_cols(rows))
    point_columns = build_point_columns(rows, block["start_col"], end_col, type_config)

    records = []
    if not point_columns:
        return records

    # 前 4 行通常是标题、点名、空行、表头；数据从第 5 行开始，Python 下标为 4。
    for r in range(4, len(rows)):
        date = parse_date(cell(rows, r, block["start_col"]))
        if not date:
            continue

        for item in point_columns:
            value = parse_number(cell(rows, r, item["col"]))
            if value is None:
                continue
            records.append({
                "type": type_key,
                "date": date,
                "point": item["point"],
                "value": value,
            })

    return records


def extract_from_summary_sheet(workbook, input_file: str):
    """从成果统计表提取地表偏移和沉降累计值。"""
    sheet_name = CONFIG["summary_sheet_name"]

    if sheet_name not in workbook.sheetnames:
        candidates = [name for name in workbook.sheetnames if "成果统计" in normalize_text(name)]
        if not candidates:
            raise RuntimeError(f"未找到成果统计表：{input_file}")
        sheet_name = candidates[0]

    ws = workbook[sheet_name]
    rows = sheet_to_rows(ws)

    surface_blocks = find_block_starts(rows, TYPES["surface"])
    settlement_blocks = find_block_starts(rows, TYPES["settlement"])
    all_blocks = sorted(surface_blocks + settlement_blocks, key=lambda item: item["start_col"])

    return {
        "surface": [
            record
            for block in surface_blocks
            for record in extract_block_records(rows, block, all_blocks, "surface")
        ],
        "settlement": [
            record
            for block in settlement_blocks
            for record in extract_block_records(rows, block, all_blocks, "settlement")
        ],
    }


def merge_records(record_groups):
    """
    合并记录为录入模板二维表。

    输入记录：
    {
        "date": "2025-05-07",
        "point": "BP-1",
        "value": 0.0
    }
    """
    date_map = defaultdict(dict)
    points = set()

    for group in record_groups:
        for record in group:
            date = record["date"]
            point = record["point"]
            value = record["value"]
            points.add(point)
            date_map[date][point] = value

    sorted_points = sorted(points, key=lambda item: [str(part) for part in re.split(r"(\d+)", item)])
    sorted_dates = sorted(date_map.keys())

    rows = [["监测日期", *sorted_points, "备注"]]
    for date in sorted_dates:
        rows.append([
            date,
            *[date_map[date].get(point, "") for point in sorted_points],
            "",
        ])

    return rows


def write_workbook(output_path: Path, sheet_name: str, rows):
    """写出 Excel。"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    for row in rows:
        ws.append(row)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def split_extracted_by_slope_rule(input_path: Path, extracted):
    base = input_path.stem
    rule = next(
        (item for item in SLOPE_SPLIT_RULES if item["source_base_name"] in base),
        None,
    )
    if not rule:
        return [{"base": base, "extracted": extracted, "split_from": ""}]

    split_points = {normalize_point_key(point) for point in rule["points"]}
    parent = {"surface": [], "settlement": []}
    target = {"surface": [], "settlement": []}

    for type_key in ("surface", "settlement"):
        for record in extracted.get(type_key, []):
            if normalize_point_key(record["point"]) in split_points:
                target[type_key].append(record)
            else:
                parent[type_key].append(record)

    return [
        {"base": base, "extracted": parent, "split_from": ""},
        {"base": rule["target_base_name"], "extracted": target, "split_from": base},
    ]


def convert():
    output_dir = Path(CONFIG["output_dir"])
    output_dir.mkdir(parents=True, exist_ok=True)

    all_surface = []
    all_settlement = []
    detail_rows = [["源文件", "地表偏移记录数", "沉降记录数"]]

    for input_file in CONFIG["input_files"]:
        input_path = Path(input_file)
        if not input_path.exists():
            print(f"跳过，文件不存在：{input_path}")
            continue

        workbook = load_workbook(input_path, data_only=True)
        extracted = extract_from_summary_sheet(workbook, str(input_path))

        for split_item in split_extracted_by_slope_rule(input_path, extracted):
            surface_records = split_item["extracted"]["surface"]
            settlement_records = split_item["extracted"]["settlement"]
            if not surface_records and not settlement_records:
                continue

            all_surface.append(surface_records)
            all_settlement.append(settlement_records)

            detail_rows.append([
                f"{input_path} -> {split_item['base']}" if split_item["split_from"] else str(input_path),
                len(surface_records),
                len(settlement_records),
            ])

            base = safe_name(split_item["base"])

            if surface_records:
                write_workbook(
                    output_dir / f"{base}_地表偏移数据_系统录入模板.xlsx",
                    TYPES["surface"]["output_sheet_name"],
                    merge_records([surface_records]),
                )

            if settlement_records:
                write_workbook(
                    output_dir / f"{base}_沉降数据_系统录入模板.xlsx",
                    TYPES["settlement"]["output_sheet_name"],
                    merge_records([settlement_records]),
                )

    if all_surface:
        write_workbook(
            output_dir / "三标_地表偏移数据_合并系统录入模板.xlsx",
            TYPES["surface"]["output_sheet_name"],
            merge_records(all_surface),
        )

    if all_settlement:
        write_workbook(
            output_dir / "三标_沉降数据_合并系统录入模板.xlsx",
            TYPES["settlement"]["output_sheet_name"],
            merge_records(all_settlement),
        )

    write_workbook(
        output_dir / "转换明细.xlsx",
        "转换明细",
        detail_rows,
    )

    print("转换完成")
    print(f"输出目录：{output_dir}")
    for file in sorted(output_dir.glob("*.xlsx")):
        print(f"- {file.name}")


if __name__ == "__main__":
    convert()
